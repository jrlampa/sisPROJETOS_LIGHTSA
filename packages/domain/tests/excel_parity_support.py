from __future__ import annotations

import math
import re
import shutil
import tempfile
import time
import unicodedata
from dataclasses import dataclass
from pathlib import Path

import openpyxl
import pytest

from packages.domain.intake.models import TipoProjeto
from packages.domain.tracao.legacy_engine.ponto_blocks import (
    BTTraversalInput,
    BTZeroTraversalInput,
    MTTraversalInput,
    RamaisTraversalInput,
)


def norm_text(value: object) -> str:
    text = str(value or "")
    text = unicodedata.normalize("NFKD", text)
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    return re.sub(r"\s+", " ", text.strip().upper())


def to_float(value: object, default: float = 0.0) -> float:
    if value is None:
        return default
    if isinstance(value, (int, float)):
        return float(value)

    text = str(value).strip()
    if not text or re.fullmatch(r"[A-Za-z]{1,4}\d{1,6}", text):
        return default

    text = text.replace(" ", "").replace("%", "").replace("V", "").replace("DAN", "")
    text = text.replace("°", "").replace("=", "")
    if not re.fullmatch(r"[-+]?\d+(?:[\.,]\d+)?", text):
        return default
    return float(text.replace(",", "."))


def load_workbook_or_skip(path: Path):
    if not path.exists():
        pytest.skip(f"Workbook nao encontrado: {path}")

    temp_dir = tempfile.TemporaryDirectory(prefix="excel-parity-")
    temp_copy = Path(temp_dir.name) / path.name

    copy_error: PermissionError | None = None
    for _ in range(3):
        try:
            shutil.copy2(path, temp_copy)
            break
        except PermissionError as exc:
            copy_error = exc
            time.sleep(0.2)
    else:
        temp_dir.cleanup()
        pytest.skip(
            "Nao foi possivel criar copia temporaria do workbook. "
            f"Feche o arquivo no Excel ou verifique permissao de leitura: {copy_error}"
        )

    try:
        return openpyxl.load_workbook(temp_copy, data_only=True, read_only=True), temp_dir
    except PermissionError as exc:
        temp_dir.cleanup()
        pytest.skip(f"Workbook indisponivel para leitura da copia temporaria: {exc}")


def iter_cells(ws, max_rows: int = 350, max_cols: int = 220):
    for row in ws.iter_rows(
        min_row=1,
        max_row=min(ws.max_row, max_rows),
        min_col=1,
        max_col=min(ws.max_column, max_cols),
        values_only=False,
    ):
        for cell in row:
            yield cell


def find_sheet_with_keywords(wb, *keywords: str):
    normalized = [norm_text(keyword) for keyword in keywords]
    for ws in wb.worksheets:
        haystack = "\n".join(
            norm_text(cell.value) for cell in iter_cells(ws) if isinstance(cell.value, str)
        )
        if all(keyword in haystack for keyword in normalized):
            return ws
    return None


def find_cell(ws, keyword: str):
    needle = norm_text(keyword)
    for cell in iter_cells(ws):
        if needle in norm_text(cell.value):
            return cell
    return None


def condutor_from_text(nome_raw: str, corrente_a: float) -> dict[str, float | str]:
    nome = (nome_raw or "").strip() or "240 Al - Arm"
    normalized = norm_text(nome)
    ampacidade_min = max(corrente_a * 1.1, 80.0)
    if "240" in normalized:
        return {
            "nome": nome,
            "resistencia_ohm_km": 0.15,
            "ampacidade_a": max(426.0, ampacidade_min),
        }
    if "70" in normalized:
        return {
            "nome": nome,
            "resistencia_ohm_km": 0.45,
            "ampacidade_a": max(140.0, ampacidade_min),
        }
    if "16" in normalized:
        return {"nome": nome, "resistencia_ohm_km": 1.15, "ampacidade_a": max(80.0, ampacidade_min)}
    return {"nome": nome, "resistencia_ohm_km": 0.45, "ampacidade_a": max(140.0, ampacidade_min)}


@dataclass
class CqtExtract:
    payload: dict
    expected_queda_total_percent: float


def extract_cqt(ws) -> CqtExtract:
    header = find_cell(ws, "Trecho do Circuito")
    if not header:
        raise AssertionError("Nao foi possivel localizar cabecalho principal de CQT no workbook.")

    col_trecho = header.column
    row0 = header.row
    col_fases_cell = find_cell(ws, "Nº de fases do trecho")
    col_secao_cell = find_cell(ws, "Seção")
    if not (col_fases_cell and col_secao_cell):
        raise AssertionError("Nao foi possivel mapear colunas essenciais de CQT no workbook.")

    col_fases = col_fases_cell.column
    col_secao = col_secao_cell.column
    col_comp = col_secao + 4
    col_carga_kva = col_secao + 36
    col_res_ajustada = col_secao + 11
    col_qdt_trecho = col_secao + 39
    col_qdt_acum = col_secao + 40

    trechos = []
    ultima_linha_trecho = None
    regex_trecho = re.compile(r"^(TR(?:AF\.?\d*)?|P-\d+|RAMAL)$", flags=re.IGNORECASE)
    for row in range(row0 + 3, min(ws.max_row, row0 + 80)):
        trecho_nome = str(ws.cell(row=row, column=col_trecho).value or "").strip()
        if not regex_trecho.match(trecho_nome):
            continue

        carga_kva = to_float(ws.cell(row=row, column=col_carga_kva).value, 0.0)
        comprimento = to_float(ws.cell(row=row, column=col_comp).value, 0.0)
        if carga_kva <= 0 or comprimento <= 0:
            continue

        fases = int(round(to_float(ws.cell(row=row, column=col_fases).value, 3)))
        secao = str(ws.cell(row=row, column=col_secao).value or "").strip()
        resistencia_ohm_km = to_float(ws.cell(row=row, column=col_res_ajustada).value, 0.0)
        qdt_trecho_excel = to_float(ws.cell(row=row, column=col_qdt_trecho).value, 0.0)
        tipo_rede = "ramal" if norm_text(trecho_nome) == "RAMAL" else "rede"
        tensao_v = 220.0
        corrente = (
            (carga_kva * 1000) / tensao_v
            if tipo_rede == "ramal"
            else (carga_kva * 1000) / (math.sqrt(3) * tensao_v)
        )
        condutor = condutor_from_text(secao, corrente)
        if resistencia_ohm_km > 0:
            condutor["resistencia_ohm_km"] = resistencia_ohm_km

        trechos.append(
            {
                "nome": trecho_nome
                if len(trecho_nome) >= 3
                else f"{trecho_nome}-{len(trechos) + 1}",
                "tipo_rede": tipo_rede,
                "fases": 1 if tipo_rede == "ramal" else max(1, min(3, fases)),
                "comprimento_m": comprimento,
                "corrente_a": corrente,
                "tensao_nominal_v": tensao_v,
                "ordem_no_circuito": len(trechos) + 1,
                "consumidores_montante": 1,
                "consumidores_jusante": 1,
                "fases_montante": 1 if tipo_rede == "ramal" else 3,
                "fases_jusante": 1 if tipo_rede == "ramal" else 3,
                "condutor": condutor,
                "queda_tensao_percent_legacy": qdt_trecho_excel,
            }
        )
        ultima_linha_trecho = row

    if not trechos or ultima_linha_trecho is None:
        raise AssertionError("Nenhum trecho CQT foi extraido do workbook.")

    primeira_linha = trechos[0]["ordem_no_circuito"] + row0 + 2
    qdt_primeiro_trecho = to_float(
        ws.cell(row=primeira_linha, column=col_qdt_trecho).value, default=float("nan")
    )
    qdt_primeiro_acumulado = to_float(
        ws.cell(row=primeira_linha, column=col_qdt_acum).value, default=float("nan")
    )
    if math.isnan(qdt_primeiro_trecho) or math.isnan(qdt_primeiro_acumulado):
        raise AssertionError(
            "Nao foi possivel extrair o acumulado inicial de QDT para calcular queda base."
        )

    expected_total = to_float(
        ws.cell(row=ultima_linha_trecho, column=col_qdt_acum).value, default=float("nan")
    )
    if math.isnan(expected_total):
        expected_total = to_float(
            ws.cell(row=ultima_linha_trecho, column=col_qdt_acum + 7).value, default=float("nan")
        )
    if math.isnan(expected_total):
        raise AssertionError(
            "Nao foi possivel extrair valor numerico de queda total acumulada no workbook."
        )

    return CqtExtract(
        payload={
            "tipo_projeto": TipoProjeto.ROBUSTEZ_BT.value,
            "recebeu_leitura_trafo_maxima": True,
            "corrente_trafo_a": 112.0,
            "carga_maxima_transformador_kva": 68.0,
            "centro_carga": {
                "nome": "CC-EXCEL",
                "queda_base_percent": max(0.0, qdt_primeiro_acumulado - qdt_primeiro_trecho),
                "transformador": {
                    "descricao": "Trafo extraido de planilha",
                    "potencia_nominal_kva": 112.5,
                    "carga_maxima_lida_kva": 68.0,
                    "corrente_lida_a": 112.0,
                },
                "trechos": trechos,
            },
        },
        expected_queda_total_percent=expected_total,
    )


@dataclass
class TracaoExtract:
    mt1: list[MTTraversalInput]
    mt2: list[MTTraversalInput]
    bt: list[BTTraversalInput]
    btz: list[BTZeroTraversalInput]
    ral: list[RamaisTraversalInput]
    tipo_poste: str
    modelo_poste: str
    expected_total_tracao: float
    expected_total_angulo: float


def find_row_by_label(ws, anchor_row: int, label: str, max_delta: int = 25) -> int | None:
    target = norm_text(label)
    for row in range(anchor_row, min(ws.max_row, anchor_row + max_delta)):
        for col in range(1, 25):
            if target in norm_text(ws.cell(row=row, column=col).value):
                return row
    return None


def parse_total_tracao_from_sheet(ws) -> tuple[float, float]:
    precise_total_row = find_row_by_label(ws, 1, "RESULTANTE - TOTAL", max_delta=ws.max_row)
    if precise_total_row is not None:
        total = to_float(ws.cell(row=precise_total_row, column=3).value, default=float("nan"))
        angulo = to_float(ws.cell(row=precise_total_row + 1, column=3).value, default=float("nan"))
        if not math.isnan(total) and not math.isnan(angulo):
            return total, angulo

    regex = re.compile(r"TRACAO\s+TOTAL\s*:\s*([\d\.,]+)\s*DAN\s*([\d\.,]+)")
    for cell in iter_cells(ws):
        match = regex.search(norm_text(cell.value))
        if match:
            return to_float(match.group(1)), to_float(match.group(2))
    raise AssertionError("Nao foi possivel localizar texto 'TRACAO TOTAL' na planilha de Tracao.")


def cell_value(ws, row: int | None, col: int, default: object = ""):
    return default if row is None or row < 1 or col < 1 else ws.cell(row=row, column=col).value


def find_section_anchor(ws, heading: str) -> int | None:
    target = norm_text(heading)
    exact_candidates: list[int] = []
    partial_candidates: list[int] = []
    for cell in iter_cells(ws, max_rows=400, max_cols=120):
        value = norm_text(cell.value)
        if value == target:
            exact_candidates.append(cell.row)
        elif target in value:
            partial_candidates.append(cell.row)

    for row in [*exact_candidates, *partial_candidates]:
        titles = {
            norm_text(ws.cell(row=row + 1, column=col).value)
            for col in range(1, min(ws.max_column, 120))
        }
        if "T1" in titles and ({"T2", "T3", "T4"} & titles):
            return row
    candidates = [*exact_candidates, *partial_candidates]
    return candidates[0] if candidates else None


def find_traversal_columns(ws, anchor_row: int) -> list[int]:
    trav_cols = [
        col
        for col in range(1, min(ws.max_column, 120))
        if norm_text(ws.cell(row=anchor_row + 1, column=col).value) in {"T1", "T2", "T3", "T4"}
    ]
    if len(trav_cols) >= 4:
        return trav_cols[:4]

    fallback_cols = [3, 6, 9, 12]
    has_t1 = any(
        norm_text(ws.cell(row=anchor_row + 1, column=col).value) == "T1" for col in fallback_cols
    )
    return fallback_cols if has_t1 else trav_cols[:4]


def extract_section_grid(ws, heading: str, label_map: dict[str, str]) -> list[dict[str, object]]:
    anchor_row = find_section_anchor(ws, heading)
    if anchor_row is None:
        return [{} for _ in range(4)]

    trav_cols = find_traversal_columns(ws, anchor_row)
    if len(trav_cols) < 4:
        return [{} for _ in range(4)]

    rows = {
        field: find_row_by_label(ws, anchor_row + 2, label) for field, label in label_map.items()
    }
    extracted: list[dict[str, object]] = []
    for col in trav_cols:
        extracted.append({field: cell_value(ws, row, col, "") for field, row in rows.items()})
    return extracted


def extract_mt_section(ws, heading: str) -> list[MTTraversalInput]:
    rows = extract_section_grid(
        ws,
        heading,
        {
            "tipo_rede": "Tipo de rede",
            "tipo_cabo": "Tipo de cabo",
            "vao": "Vao",
            "flecha": "Flecha",
            "angulo": "Angulo",
            "altura_poste": "Altura poste",
            "altura_ancoragem": "Altura ancoragem",
        },
    )
    return [
        MTTraversalInput(
            tipo_rede=str(row.get("tipo_rede") or "").strip(),
            tipo_cabo=str(row.get("tipo_cabo") or "").strip(),
            vao=to_float(row.get("vao"), 0.0),
            flecha=to_float(row.get("flecha"), 0.0),
            angulo=to_float(row.get("angulo"), 0.0),
            altura_poste=to_float(row.get("altura_poste"), 0.0),
            altura_ancoragem=to_float(row.get("altura_ancoragem"), 0.0),
        )
        for row in rows
    ]


def extract_tracao(ws) -> TracaoExtract:
    tipo_poste_cell = find_cell(ws, "Tipo do Poste")
    modelo_poste_cell = find_cell(ws, "Modelo do Poste")
    if not (tipo_poste_cell and modelo_poste_cell):
        raise AssertionError(
            "Nao foi possivel localizar bloco fixo de DADOS DO POSTE na planilha de Tracao."
        )

    mt1 = extract_mt_section(ws, "MT - 1")
    mt2 = extract_mt_section(ws, "MT - 2")
    bt = [
        BTTraversalInput(
            tipo_rede=item.tipo_rede,
            tipo_cabo=item.tipo_cabo,
            vao=item.vao,
            flecha=item.flecha,
            angulo=item.angulo,
            altura_poste=item.altura_poste,
            altura_ancoragem=item.altura_ancoragem,
        )
        for item in extract_mt_section(ws, "BT")
    ]
    btz = [
        BTZeroTraversalInput(
            qtd_ligacoes=to_float(row.get("qtd_ligacoes"), 0.0),
            vao=to_float(row.get("vao"), 0.0),
            flecha=to_float(row.get("flecha"), 0.0),
            angulo=to_float(row.get("angulo"), 0.0),
            altura_poste=to_float(row.get("altura_poste"), 0.0),
            altura_ancoragem=to_float(row.get("altura_ancoragem"), 0.0),
        )
        for row in extract_section_grid(
            ws,
            "RAMAIS BTZERO",
            {
                "qtd_ligacoes": "Quantidade de ligacoes",
                "vao": "Vao",
                "flecha": "Flecha",
                "angulo": "Angulo",
                "altura_poste": "Altura poste",
                "altura_ancoragem": "Altura ancoragem",
            },
        )
    ]
    ral = [
        RamaisTraversalInput(
            tipo_cabo=str(row.get("tipo_cabo") or "").strip(),
            qtd_cabos=to_float(row.get("qtd_cabos"), 0.0),
            vao=to_float(row.get("vao"), 0.0),
            flecha=to_float(row.get("flecha"), 0.0),
            angulo=to_float(row.get("angulo"), 0.0),
            altura_poste=to_float(row.get("altura_poste"), 0.0),
            altura_ancoragem=to_float(row.get("altura_ancoragem"), 0.0),
        )
        for row in extract_section_grid(
            ws,
            "RAMAIS DE LIGACAO",
            {
                "tipo_cabo": "Tipo de cabo",
                "qtd_cabos": "Quantidade de cabos",
                "vao": "Vao",
                "flecha": "Flecha",
                "angulo": "Angulo",
                "altura_poste": "Altura poste",
                "altura_ancoragem": "Altura ancoragem",
            },
        )
    ]
    total_tracao, total_angulo = parse_total_tracao_from_sheet(ws)
    return TracaoExtract(
        mt1=mt1,
        mt2=mt2,
        bt=bt,
        btz=btz,
        ral=ral,
        tipo_poste=str(
            ws.cell(row=tipo_poste_cell.row, column=tipo_poste_cell.column + 1).value or ""
        ).strip(),
        modelo_poste=str(
            ws.cell(row=modelo_poste_cell.row, column=modelo_poste_cell.column + 1).value or ""
        ).strip(),
        expected_total_tracao=total_tracao,
        expected_total_angulo=total_angulo,
    )
