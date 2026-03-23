"""Paridade de calculo entre backend e workbook Excel legado LIGHT.

Este teste e de integracao com arquivo real (.xlsm).
Ele tenta extrair entradas e saidas calculadas diretamente do workbook,
executa os motores de CQT e Tracao e compara os resultados com tolerancia
maxima de 0.01% (relativa).
"""

from __future__ import annotations

import math
import re
import shutil
import tempfile
import time
import unicodedata
from dataclasses import dataclass
from pathlib import Path

import openpyxl
import pytest

from packages.domain.cqt.models import CQTAnalise
from packages.domain.intake.models import TipoProjeto
from packages.domain.tracao.legacy_engine.ponto_blocks import (
    BTTraversalInput,
    BTZeroTraversalInput,
    MTTraversalInput,
    RamaisTraversalInput,
    calcular_polo,
)

EXCEL_PATH = Path(
    r"C:\Users\jonat\OneDrive - IM3 Brasil\LIGHT\PROJETOS\ZNA41723 - ROBUSTEZ DE BT\QDT_ZNA_41723_ATUAL.xlsm"
)
MAX_DIVERGENCIA_PERCENT = 0.01


def _norm(value: object) -> str:
    text = str(value or "")
    text = unicodedata.normalize("NFKD", text)
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    return re.sub(r"\s+", " ", text.strip().upper())


def _to_float(value: object, default: float = 0.0) -> float:
    if value is None:
        return default
    if isinstance(value, (int, float)):
        return float(value)

    txt = str(value).strip()
    if not txt:
        return default

    txt = txt.replace(" ", "").replace("%", "").replace("V", "").replace("DAN", "")
    txt = txt.replace("°", "").replace("=", "")
    txt = txt.replace(".", "").replace(",", ".")

    m = re.search(r"[-+]?\d+(?:\.\d+)?", txt)
    if not m:
        return default
    return float(m.group(0))


def _assert_close_rel_percent(actual: float, expected: float, label: str) -> None:
    denom = max(abs(expected), 1e-12)
    diff_percent = (abs(actual - expected) / denom) * 100
    assert diff_percent <= MAX_DIVERGENCIA_PERCENT, (
        f"Divergencia acima da tolerancia em {label}: "
        f"actual={actual:.10f}, expected={expected:.10f}, diff%={diff_percent:.10f}, "
        f"tol%={MAX_DIVERGENCIA_PERCENT}"
    )


def _load_workbook_or_skip(path: Path):
    if openpyxl is None:
        pytest.skip("Dependencia openpyxl nao instalada no ambiente.")

    if not path.exists():
        pytest.skip(f"Workbook nao encontrado: {path}")

    temp_dir = tempfile.TemporaryDirectory(prefix="excel-parity-")
    temp_copy = Path(temp_dir.name) / path.name

    copy_error: PermissionError | None = None
    for _ in range(3):
        try:
            shutil.copy2(path, temp_copy)
            break
        except PermissionError as exc:
            copy_error = exc
            time.sleep(0.2)
    else:
        temp_dir.cleanup()
        pytest.skip(
            "Nao foi possivel criar copia temporaria do workbook. "
            f"Feche o arquivo no Excel ou verifique permissao de leitura: {copy_error}"
        )

    try:
        wb = openpyxl.load_workbook(temp_copy, data_only=True, read_only=True)
        return wb, temp_dir
    except PermissionError as exc:
        temp_dir.cleanup()
        pytest.skip(f"Workbook indisponivel para leitura da copia temporaria: {exc}")


def _iter_cells(ws, max_rows: int = 350, max_cols: int = 220):
    for row in ws.iter_rows(
        min_row=1,
        max_row=min(ws.max_row, max_rows),
        min_col=1,
        max_col=min(ws.max_column, max_cols),
        values_only=False,
    ):
        for cell in row:
            yield cell


def _find_sheet_with_keywords(wb, *keywords: str):
    normalized = [_norm(k) for k in keywords]
    for ws in wb.worksheets:
        text_pool = []
        for cell in _iter_cells(ws):
            if isinstance(cell.value, str):
                text_pool.append(_norm(cell.value))
        haystack = "\n".join(text_pool)
        if all(key in haystack for key in normalized):
            return ws
    return None


def _find_cell(ws, keyword: str):
    needle = _norm(keyword)
    for cell in _iter_cells(ws):
        if needle in _norm(cell.value):
            return cell
    return None


def _first_numeric_right(ws, row: int, col: int, search_width: int = 10) -> float | None:
    for c in range(col + 1, col + search_width + 1):
        value = ws.cell(row=row, column=c).value
        parsed = _to_float(value, default=float("nan"))
        if not math.isnan(parsed):
            return parsed
    return None


def _condutor_from_text(nome_raw: str, corrente_a: float) -> dict[str, float | str]:
    nome = (nome_raw or "").strip() or "240 Al - Arm"
    up = _norm(nome)
    ampacidade_min = max(corrente_a * 1.1, 80.0)

    if "240" in up:
        return {
            "nome": nome,
            "resistencia_ohm_km": 0.15,
            "ampacidade_a": max(426.0, ampacidade_min),
        }
    if "70" in up:
        return {
            "nome": nome,
            "resistencia_ohm_km": 0.45,
            "ampacidade_a": max(140.0, ampacidade_min),
        }
    if "16" in up:
        return {"nome": nome, "resistencia_ohm_km": 1.15, "ampacidade_a": max(80.0, ampacidade_min)}
    return {"nome": nome, "resistencia_ohm_km": 0.45, "ampacidade_a": max(140.0, ampacidade_min)}


@dataclass
class CqtExtract:
    payload: dict
    expected_queda_total_percent: float


def _extract_cqt(ws) -> CqtExtract:
    header = _find_cell(ws, "Trecho do Circuito")
    if not header:
        raise AssertionError("Nao foi possivel localizar cabecalho principal de CQT no workbook.")

    row0 = header.row
    col_trecho = header.column

    col_fases = _find_cell(ws, "Nº de fases do trecho")
    col_secao = _find_cell(ws, "Seção")
    col_corr = _find_cell(ws, "Corrente do cabo para a condição")

    if not (col_fases and col_secao and col_corr):
        raise AssertionError("Nao foi possivel mapear colunas essenciais de CQT no workbook.")

    col_fases = col_fases.column
    col_secao = col_secao.column
    col_corr = col_corr.column
    col_comp = col_corr + 2

    trechos = []
    regex_trecho = re.compile(r"^(TR|P-\d+|RAMAL)$", flags=re.IGNORECASE)

    for r in range(row0 + 3, min(ws.max_row, row0 + 80)):
        trecho_nome = str(ws.cell(row=r, column=col_trecho).value or "").strip()
        if not regex_trecho.match(trecho_nome):
            continue

        fases = int(round(_to_float(ws.cell(row=r, column=col_fases).value, 3)))
        corrente = _to_float(ws.cell(row=r, column=col_corr).value, 0.0)
        comprimento = _to_float(ws.cell(row=r, column=col_comp).value, 0.0)
        secao = str(ws.cell(row=r, column=col_secao).value or "").strip()

        if corrente <= 0 or comprimento <= 0:
            continue

        tipo_rede = "ramal" if _norm(trecho_nome) == "RAMAL" else "rede"
        nome_trecho = trecho_nome if len(trecho_nome) >= 3 else f"{trecho_nome}-{len(trechos) + 1}"
        trechos.append(
            {
                "nome": nome_trecho,
                "tipo_rede": tipo_rede,
                "fases": 1 if tipo_rede == "ramal" else max(1, min(3, fases)),
                "comprimento_m": comprimento,
                "corrente_a": corrente,
                "tensao_nominal_v": 220.0,
                "ordem_no_circuito": len(trechos) + 1,
                "consumidores_montante": 1,
                "consumidores_jusante": 1,
                "fases_montante": 1 if tipo_rede == "ramal" else 3,
                "fases_jusante": 1 if tipo_rede == "ramal" else 3,
                "condutor": _condutor_from_text(secao, corrente),
            }
        )

    if not trechos:
        raise AssertionError("Nenhum trecho CQT foi extraido do workbook.")

    label_total = _find_cell(ws, "ΔV (%) no Cabo de BT (total ACUMULADO)") or _find_cell(
        ws, "DV (%) no Cabo de BT (total ACUMULADO)"
    )
    if not label_total:
        raise AssertionError(
            "Nao foi possivel localizar a referencia de queda total acumulada no workbook."
        )

    expected_total = _first_numeric_right(ws, label_total.row, label_total.column)
    if expected_total is None:
        raise AssertionError(
            "Nao foi possivel extrair valor numerico de queda total acumulada no workbook."
        )

    payload = {
        "tipo_projeto": TipoProjeto.ROBUSTEZ_BT.value,
        "recebeu_leitura_trafo_maxima": True,
        "corrente_trafo_a": 112.0,
        "carga_maxima_transformador_kva": 68.0,
        "centro_carga": {
            "nome": "CC-EXCEL",
            "transformador": {
                "descricao": "Trafo extraido de planilha",
                "potencia_nominal_kva": 112.5,
                "carga_maxima_lida_kva": 68.0,
                "corrente_lida_a": 112.0,
            },
            "trechos": trechos,
        },
    }

    return CqtExtract(payload=payload, expected_queda_total_percent=expected_total)


@dataclass
class TracaoExtract:
    mt1: list[MTTraversalInput]
    mt2: list[MTTraversalInput]
    bt: list[BTTraversalInput]
    btz: list[BTZeroTraversalInput]
    ral: list[RamaisTraversalInput]
    tipo_poste: str
    modelo_poste: str
    expected_total_tracao: float
    expected_total_angulo: float


def _parse_total_tracao_from_sheet(ws) -> tuple[float, float]:
    regex = re.compile(r"TRACAO\s+TOTAL\s*:\s*([\d\.,]+)\s*DAN\s*([\d\.,]+)")
    for cell in _iter_cells(ws):
        text = str(cell.value or "")
        norm = _norm(text)
        m = regex.search(norm)
        if m:
            return _to_float(m.group(1)), _to_float(m.group(2))
    raise AssertionError("Nao foi possivel localizar texto 'TRACAO TOTAL' na planilha de Tracao.")


def _find_row_by_label(ws, anchor_row: int, label: str, max_delta: int = 25) -> int | None:
    target = _norm(label)
    for r in range(anchor_row, min(ws.max_row, anchor_row + max_delta)):
        v = ws.cell(row=r, column=1).value
        if target in _norm(v):
            return r
    return None


def _extract_mt_section(ws, heading: str) -> list[MTTraversalInput]:
    anchor = _find_cell(ws, heading)
    if not anchor:
        return [MTTraversalInput() for _ in range(4)]

    row = anchor.row
    trav_cols = []
    for c in range(1, min(ws.max_column, 80)):
        if _norm(ws.cell(row=row + 1, column=c).value) in {"T1", "T2", "T3", "T4"}:
            trav_cols.append(c)
    if len(trav_cols) < 4:
        return [MTTraversalInput() for _ in range(4)]

    labels = {
        "tipo_rede": _find_row_by_label(ws, row + 2, "Tipo de rede"),
        "tipo_cabo": _find_row_by_label(ws, row + 2, "Tipo de cabo"),
        "vao": _find_row_by_label(ws, row + 2, "Vao"),
        "flecha": _find_row_by_label(ws, row + 2, "Flecha"),
        "angulo": _find_row_by_label(ws, row + 2, "Angulo"),
        "altura_poste": _find_row_by_label(ws, row + 2, "Altura poste"),
        "altura_ancoragem": _find_row_by_label(ws, row + 2, "Altura ancoragem"),
    }

    out: list[MTTraversalInput] = []
    for col in trav_cols[:4]:
        out.append(
            MTTraversalInput(
                tipo_rede=str(
                    ws.cell(row=labels["tipo_rede"] or 0, column=col).value or ""
                ).strip(),
                tipo_cabo=str(
                    ws.cell(row=labels["tipo_cabo"] or 0, column=col).value or ""
                ).strip(),
                vao=_to_float(ws.cell(row=labels["vao"] or 0, column=col).value, 0.0),
                flecha=_to_float(ws.cell(row=labels["flecha"] or 0, column=col).value, 0.0),
                angulo=_to_float(ws.cell(row=labels["angulo"] or 0, column=col).value, 0.0),
                altura_poste=_to_float(
                    ws.cell(row=labels["altura_poste"] or 0, column=col).value, 0.0
                ),
                altura_ancoragem=_to_float(
                    ws.cell(row=labels["altura_ancoragem"] or 0, column=col).value, 0.0
                ),
            )
        )
    return out


def _extract_tracao(ws) -> TracaoExtract:
    tipo_poste_cell = _find_cell(ws, "Tipo do Poste")
    modelo_poste_cell = _find_cell(ws, "Modelo do Poste")
    if not (tipo_poste_cell and modelo_poste_cell):
        raise AssertionError(
            "Nao foi possivel localizar bloco fixo de DADOS DO POSTE na planilha de Tracao."
        )

    tipo_poste = str(
        ws.cell(row=tipo_poste_cell.row, column=tipo_poste_cell.column + 1).value or ""
    ).strip()
    modelo_poste = str(
        ws.cell(row=modelo_poste_cell.row, column=modelo_poste_cell.column + 1).value or ""
    ).strip()

    mt1 = _extract_mt_section(ws, "MT - 1")
    mt2 = _extract_mt_section(ws, "MT - 2")

    bt_mt = _extract_mt_section(ws, "BT")
    bt = [
        BTTraversalInput(
            tipo_rede=t.tipo_rede,
            tipo_cabo=t.tipo_cabo,
            vao=t.vao,
            flecha=t.flecha,
            angulo=t.angulo,
            altura_poste=t.altura_poste,
            altura_ancoragem=t.altura_ancoragem,
        )
        for t in bt_mt
    ]

    btz = [BTZeroTraversalInput() for _ in range(4)]
    ral = [RamaisTraversalInput() for _ in range(4)]

    total_tracao, total_angulo = _parse_total_tracao_from_sheet(ws)

    return TracaoExtract(
        mt1=mt1,
        mt2=mt2,
        bt=bt,
        btz=btz,
        ral=ral,
        tipo_poste=tipo_poste,
        modelo_poste=modelo_poste,
        expected_total_tracao=total_tracao,
        expected_total_angulo=total_angulo,
    )


def test_excel_parity_cqt_e_tracao() -> None:
    wb, temp_dir = _load_workbook_or_skip(EXCEL_PATH)

    try:
        ws_cqt = _find_sheet_with_keywords(wb, "Trecho do Circuito", "Queda")
        if ws_cqt is None:
            raise AssertionError("Nao foi encontrada aba de CQT na planilha fornecida.")

        cqt = _extract_cqt(ws_cqt)
        analise = CQTAnalise.model_validate(cqt.payload)
        _assert_close_rel_percent(
            actual=analise.centro_carga.queda_total_percent,
            expected=cqt.expected_queda_total_percent,
            label="CQT.queda_total_percent",
        )

        ws_tracao = _find_sheet_with_keywords(wb, "TRACAO TOTAL")
        if ws_tracao is None:
            raise AssertionError("Nao foi encontrada aba de Tracao na planilha fornecida.")

        tracao = _extract_tracao(ws_tracao)
        out = calcular_polo(
            mt1_inputs=tracao.mt1,
            mt2_inputs=tracao.mt2,
            bt_inputs=tracao.bt,
            btz_inputs=tracao.btz,
            ral_inputs=tracao.ral,
            tipo_poste=tracao.tipo_poste,
            modelo_poste=tracao.modelo_poste,
        )

        _assert_close_rel_percent(
            actual=out.total_tracao,
            expected=tracao.expected_total_tracao,
            label="TRACAO.total_tracao",
        )
        _assert_close_rel_percent(
            actual=out.total_angulo,
            expected=tracao.expected_total_angulo,
            label="TRACAO.total_angulo",
        )

        # Momentos/resultantes visiveis para QA eletrica (componentes vetoriais do total)
        expected_mx = tracao.expected_total_tracao * math.cos(
            math.radians(tracao.expected_total_angulo)
        )
        expected_my = tracao.expected_total_tracao * math.sin(
            math.radians(tracao.expected_total_angulo)
        )
        actual_mx = out.total_tracao * math.cos(math.radians(out.total_angulo))
        actual_my = out.total_tracao * math.sin(math.radians(out.total_angulo))

        _assert_close_rel_percent(actual=actual_mx, expected=expected_mx, label="TRACAO.momento_x")
        _assert_close_rel_percent(actual=actual_my, expected=expected_my, label="TRACAO.momento_y")
    finally:
        wb.close()
        temp_dir.cleanup()
