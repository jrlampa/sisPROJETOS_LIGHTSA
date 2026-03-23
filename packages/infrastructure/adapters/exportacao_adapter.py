"""Adaptador de infraestrutura para gerar artefatos fisicos de exportacao."""

from __future__ import annotations

import tempfile
import zipfile
from pathlib import Path
from uuid import UUID

from openpyxl import Workbook
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas


class ExportadorFicheiros:
    """Gera ficheiros Excel, PDF e ZIP em diretório temporário."""

    def __init__(self, base_dir: str | None = None) -> None:
        self._base_dir = (
            Path(base_dir) if base_dir else Path(tempfile.mkdtemp(prefix="sisprojetos_export_"))
        )
        self._base_dir.mkdir(parents=True, exist_ok=True)

    def gerar_excel_padrao(
        self, projeto_id: UUID, dados_cqt: dict, dados_tracao: list[dict]
    ) -> str:
        """Cria um esqueleto de planilha para o pacote final de exportacao."""
        caminho = self._base_dir / f"{projeto_id}_relatorio_tecnico.xlsx"

        wb = Workbook()
        ws = wb.active
        ws.title = "Resumo"
        ws["A1"] = "Projeto ID"
        ws["B1"] = str(projeto_id)
        ws["A2"] = "Tipo CQT"
        ws["B2"] = dados_cqt.get("tipo_projeto", "N/A")
        ws["A3"] = "Total de Postes em Tracao"
        ws["B3"] = len(dados_tracao)
        ws["A5"] = "Estado"
        ws["B5"] = "Percentual"

        linha = 6
        for item in dados_tracao:
            ws[f"A{linha}"] = item.get("estado_mecanico", "N/A")
            ws[f"B{linha}"] = item.get("percentual_carregamento", 0.0)
            linha += 1

        wb.save(caminho)
        return str(caminho)

    def gerar_pdf_tecnico(self, projeto_id: UUID) -> str:
        """Cria um PDF simples com metadados basicos do projeto."""
        caminho = self._base_dir / f"{projeto_id}_relatorio_tecnico.pdf"

        doc = canvas.Canvas(str(caminho), pagesize=A4)
        doc.setTitle("Relatorio Tecnico")
        doc.drawString(72, 800, "sisPROJETOS LIGHT S.A. - Relatorio Tecnico")
        doc.drawString(72, 780, f"Projeto ID: {projeto_id}")
        doc.drawString(72, 760, "Documento gerado automaticamente na Onda 5.")
        doc.showPage()
        doc.save()

        return str(caminho)

    def gerar_pacote_zip(self, projeto_id: UUID, lista_caminhos: list[str]) -> str:
        """Compacta os ficheiros informados em um ZIP final de entrega."""
        caminho_zip = self._base_dir / f"{projeto_id}_pacote_final.zip"

        with zipfile.ZipFile(
            caminho_zip, mode="w", compression=zipfile.ZIP_DEFLATED
        ) as arquivo_zip:
            for caminho in lista_caminhos:
                path = Path(caminho)
                if path.exists():
                    arquivo_zip.write(path, arcname=path.name)

        return str(caminho_zip)
