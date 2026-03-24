from __future__ import annotations

import argparse
import csv
import math
import re
import shutil
import tempfile
import time
from pathlib import Path

import openpyxl

from packages.domain.tests.excel_parity_support import extract_tracao, find_sheet_with_keywords
from packages.domain.tracao.legacy_engine.ponto_blocks import calcular_polo

PROJECT_ROOT = Path(__file__).resolve().parents[1]

DEFAULT_TARGET_DIR = Path(
    r"C:\Users\jonat\OneDrive - IM3 Brasil\LIGHT\PROJETOS\REDE CLANDESTINA - RUAS JERUSALÉM E UVA - SANTA CRUZ RJ\CALC TRAÇÃO"
)
DEFAULT_REPORT_PATH = PROJECT_ROOT / "legacy_audit_report.csv"
REL_TOL = 0.0001
POSTE_OVERLOAD_TOLERANCE = 1.05
INVALID_FORMAT_STATUS = "Ignorado - Formato Invalido"


def _load_workbook_with_copy(path: Path):
    temp_dir = tempfile.TemporaryDirectory(prefix="legacy-audit-")
    temp_copy = Path(temp_dir.name) / path.name

    copy_error: PermissionError | None = None
    for _ in range(3):
        try:
            shutil.copy2(path, temp_copy)
            break
        except PermissionError as exc:
            copy_error = exc
            time.sleep(0.2)
    else:
        temp_dir.cleanup()
        raise RuntimeError(
            f"Nao foi possivel copiar o workbook (file-lock/permissao): {copy_error}"
        )

    try:
        wb = openpyxl.load_workbook(temp_copy, data_only=True, read_only=False)
    except Exception as exc:
        temp_dir.cleanup()
        raise RuntimeError(f"Falha ao abrir workbook copiado: {exc}") from exc

    return wb, temp_dir


def _parse_poste_capacidade_dan(modelo_poste: str) -> float | None:
    text = str(modelo_poste or "").upper().replace(",", ".")
    match = re.search(r"(\d+(?:\.\d+)?)\s*DAN", text)
    return float(match.group(1)) if match else None


def _is_close(a: float, b: float) -> bool:
    return math.isclose(a, b, rel_tol=REL_TOL, abs_tol=0.0)


def _collect_entry_human_errors(entries, section: str) -> list[str]:
    errors: list[str] = []
    for idx, item in enumerate(entries, start=1):
        vao = float(getattr(item, "vao", 0.0) or 0.0)
        flecha = float(getattr(item, "flecha", 0.0) or 0.0)
        angulo = float(getattr(item, "angulo", 0.0) or 0.0)
        altura_poste = float(getattr(item, "altura_poste", 0.0) or 0.0)
        altura_anc = float(getattr(item, "altura_ancoragem", 0.0) or 0.0)

        if vao <= 0:
            continue

        if section in {"MT1", "MT2", "BT"}:
            if not str(getattr(item, "tipo_rede", "") or "").strip():
                errors.append(f"{section} T{idx}: vao preenchido sem tipo de rede")
            if not str(getattr(item, "tipo_cabo", "") or "").strip():
                errors.append(f"{section} T{idx}: vao preenchido sem tipo de cabo")

        if section == "RAL":
            if not str(getattr(item, "tipo_cabo", "") or "").strip():
                errors.append(f"RAL T{idx}: vao preenchido sem tipo de cabo")
            if float(getattr(item, "qtd_cabos", 0.0) or 0.0) <= 0:
                errors.append(f"RAL T{idx}: vao preenchido sem quantidade de cabos")

        if section == "BTZ" and float(getattr(item, "qtd_ligacoes", 0.0) or 0.0) <= 0:
            errors.append(f"BTZ T{idx}: vao preenchido sem quantidade de ligacoes")

        if flecha <= 0:
            errors.append(f"{section} T{idx}: flecha nao positiva com vao preenchido")
        if not (0 <= angulo <= 360):
            errors.append(f"{section} T{idx}: angulo fora de 0..360")
        if altura_poste > 0 and altura_anc > altura_poste:
            errors.append(f"{section} T{idx}: altura de ancoragem maior que altura do poste")

    return errors


def _is_invalid_format_exception(exc: BaseException) -> bool:
    message = str(exc).lower()
    invalid_signatures = (
        "aba de tracao nao encontrada",
        "nao foi possivel localizar",
        "nao foi possivel mapear",
        "zip file",
        "file is not a zip file",
    )
    return any(signature in message for signature in invalid_signatures)


def _audit_file(workbook_path: Path) -> dict[str, str]:
    result = {
        "Arquivo": str(workbook_path),
        "Status Paridade": "ERRO_PROCESSAMENTO",
        "Erros Humanos Detectados": "",
        "Tração Python": "",
        "Tração Excel": "",
        "Divergência": "",
    }

    wb = None
    temp_dir = None
    try:
        wb, temp_dir = _load_workbook_with_copy(workbook_path)
        ws = find_sheet_with_keywords(wb, "TRACAO TOTAL")
        if ws is None:
            raise RuntimeError("Aba de tracao nao encontrada")

        tracao_excel = extract_tracao(ws)
        tracao_python = calcular_polo(
            mt1_inputs=tracao_excel.mt1,
            mt2_inputs=tracao_excel.mt2,
            bt_inputs=tracao_excel.bt,
            btz_inputs=tracao_excel.btz,
            ral_inputs=tracao_excel.ral,
            tipo_poste=tracao_excel.tipo_poste,
            modelo_poste=tracao_excel.modelo_poste,
        )

        excel_total = float(tracao_excel.expected_total_tracao)
        excel_angle = float(tracao_excel.expected_total_angulo)
        py_total = float(tracao_python.total_tracao)
        py_angle = float(tracao_python.total_angulo)

        excel_mx = excel_total * math.cos(math.radians(excel_angle))
        excel_my = excel_total * math.sin(math.radians(excel_angle))
        py_mx = py_total * math.cos(math.radians(py_angle))
        py_my = py_total * math.sin(math.radians(py_angle))

        parity_errors: list[str] = []
        if not _is_close(py_total, excel_total):
            parity_errors.append("tracao_total")
        if not _is_close(py_angle, excel_angle):
            parity_errors.append("angulo_total")
        if not _is_close(py_mx, excel_mx):
            parity_errors.append("momento_x")
        if not _is_close(py_my, excel_my):
            parity_errors.append("momento_y")

        human_errors: list[str] = []
        capacidade_dan = _parse_poste_capacidade_dan(tracao_excel.modelo_poste)
        if capacidade_dan is None:
            human_errors.append("Modelo do poste sem capacidade em daN")
        elif excel_total > capacidade_dan * POSTE_OVERLOAD_TOLERANCE:
            human_errors.append(
                f"Tracao Excel ({excel_total:.2f}) excede capacidade do poste com tolerancia de 5% ({capacidade_dan * POSTE_OVERLOAD_TOLERANCE:.2f})"
            )

        human_errors.extend(_collect_entry_human_errors(tracao_excel.mt1, "MT1"))
        human_errors.extend(_collect_entry_human_errors(tracao_excel.mt2, "MT2"))
        human_errors.extend(_collect_entry_human_errors(tracao_excel.bt, "BT"))
        human_errors.extend(_collect_entry_human_errors(tracao_excel.btz, "BTZ"))
        human_errors.extend(_collect_entry_human_errors(tracao_excel.ral, "RAL"))

        result["Status Paridade"] = (
            "OK" if not parity_errors else f"DIVERGENTE ({', '.join(parity_errors)})"
        )
        result["Erros Humanos Detectados"] = "; ".join(human_errors) if human_errors else "Nenhum"
        result["Tração Python"] = f"{py_total:.6f}"
        result["Tração Excel"] = f"{excel_total:.6f}"

        if abs(excel_total) > 1e-12:
            diff_percent = abs(py_total - excel_total) / abs(excel_total) * 100.0
            result["Divergência"] = f"{diff_percent:.6f}%"
        else:
            result["Divergência"] = "N/A"

    except BaseException as exc:
        if _is_invalid_format_exception(exc):
            result["Status Paridade"] = INVALID_FORMAT_STATUS
            result["Erros Humanos Detectados"] = f"Formato invalido: {exc}"
        else:
            result["Erros Humanos Detectados"] = f"Falha de processamento: {exc}"
    finally:
        if wb is not None:
            wb.close()
        if temp_dir is not None:
            try:
                temp_dir.cleanup()
            except OSError:
                # Em alguns casos no Windows o handle do .xlsm ainda nao foi liberado.
                pass

    return result


def run_audit(target_dir: Path, report_path: Path) -> list[dict[str, str]]:
    files = sorted(target_dir.rglob("*.xlsm"))
    if not files:
        print(f"Nenhum arquivo .xlsm encontrado em: {target_dir}")
        return []

    rows = []
    for path in files:
        try:
            rows.append(_audit_file(path))
        except BaseException as exc:
            rows.append(
                {
                    "Arquivo": str(path),
                    "Status Paridade": "ERRO_PROCESSAMENTO",
                    "Erros Humanos Detectados": f"Falha de processamento no loop: {exc}",
                    "Tração Python": "",
                    "Tração Excel": "",
                    "Divergência": "",
                }
            )

    with report_path.open("w", newline="", encoding="utf-8") as csvfile:
        writer = csv.DictWriter(
            csvfile,
            fieldnames=[
                "Arquivo",
                "Status Paridade",
                "Erros Humanos Detectados",
                "Tração Python",
                "Tração Excel",
                "Divergência",
            ],
        )
        writer.writeheader()
        writer.writerows(rows)

    total = len(rows)
    parity_ok = sum(1 for row in rows if row["Status Paridade"] == "OK")
    parity_div = sum(1 for row in rows if row["Status Paridade"].startswith("DIVERGENTE"))
    ignored_invalid = sum(1 for row in rows if row["Status Paridade"] == INVALID_FORMAT_STATUS)
    proc_error = sum(1 for row in rows if row["Status Paridade"] == "ERRO_PROCESSAMENTO")
    human_error = sum(
        1
        for row in rows
        if row["Erros Humanos Detectados"] and row["Erros Humanos Detectados"] != "Nenhum"
    )

    print("Auditoria de legado concluida")
    print(f"- Diretorio auditado: {target_dir}")
    print(f"- Arquivos processados: {total}")
    print(f"- Paridade OK: {parity_ok}")
    print(f"- Paridade divergente: {parity_div}")
    print(f"- Ignorados por formato invalido: {ignored_invalid}")
    print(f"- Erro de processamento: {proc_error}")
    print(f"- Com erros humanos detectados: {human_error}")
    print(f"- Relatorio: {report_path}")

    return rows


def _parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Audita planilhas .xlsm legadas de tracao com resiliencia."
    )
    parser.add_argument(
        "target_dir",
        nargs="?",
        default=str(DEFAULT_TARGET_DIR),
        help="Diretorio raiz para varredura de arquivos .xlsm.",
    )
    parser.add_argument(
        "--report",
        default=str(DEFAULT_REPORT_PATH),
        help="Caminho do CSV de saida.",
    )
    return parser.parse_args()


def main() -> int:
    args = _parse_args()
    target_dir = Path(args.target_dir)
    report_path = Path(args.report)
    run_audit(target_dir=target_dir, report_path=report_path)

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
