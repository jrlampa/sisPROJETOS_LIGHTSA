from __future__ import annotations

import csv
import math
import re
import shutil
import tempfile
import time
from pathlib import Path

import openpyxl
from pydantic import ValidationError

from packages.domain.cqt.models import CQTAnalise
from packages.domain.tests.excel_parity_support import (
    extract_cqt,
    find_cell,
    norm_text,
    to_float,
)

PROJECT_ROOT = Path(__file__).resolve().parents[1]
TARGET_DIR = Path(
    r"C:\Users\jonat\OneDrive - IM3 Brasil\LIGHT\PROJETOS\REDE CLANDESTINA - RUAS JERUSALÉM E UVA - SANTA CRUZ RJ"
)
REPORT_PATH = PROJECT_ROOT / "legacy_audit_report_cqt.csv"
REL_TOL = 0.0001
QDT_LIMIT_PERCENT = 5.0


def _load_workbook_with_copy(path: Path):
    temp_dir = tempfile.TemporaryDirectory(prefix="legacy-cqt-audit-")
    temp_copy = Path(temp_dir.name) / path.name

    copy_error: PermissionError | None = None
    for _ in range(3):
        try:
            shutil.copy2(path, temp_copy)
            break
        except PermissionError as exc:
            copy_error = exc
            time.sleep(0.2)
    else:
        temp_dir.cleanup()
        raise RuntimeError(
            f"Nao foi possivel copiar o workbook (file-lock/permissao): {copy_error}"
        )

    try:
        wb = openpyxl.load_workbook(temp_copy, data_only=True, read_only=False)
    except Exception as exc:
        temp_dir.cleanup()
        raise RuntimeError(f"Falha ao abrir workbook copiado: {exc}") from exc

    return wb, temp_dir


def _is_close(a: float, b: float) -> bool:
    return math.isclose(a, b, rel_tol=REL_TOL, abs_tol=0.0)


def _extract_esforco_resultante(ws) -> float | None:
    for cell in ws.iter_rows(
        min_row=1,
        max_row=min(ws.max_row, 260),
        min_col=1,
        max_col=min(ws.max_column, 220),
        values_only=False,
    ):
        for item in cell:
            value = item.value
            if not isinstance(value, str):
                continue
            text = norm_text(value)
            if "ESFORCO" not in text and "ESFORÇO" not in text:
                continue

            row = item.row
            col = item.column
            for c in range(col + 1, min(ws.max_column, col + 12)):
                parsed = to_float(ws.cell(row=row, column=c).value, default=float("nan"))
                if not math.isnan(parsed):
                    return parsed
    return None


def _iter_cqt_worksheets(wb):
    for ws in wb.worksheets:
        has_trecho = find_cell(ws, "Trecho do Circuito") is not None
        has_queda = find_cell(ws, "Queda") is not None or find_cell(ws, "DV") is not None
        if has_trecho and has_queda:
            yield ws


def _collect_human_errors_from_sheet(ws) -> list[str]:
    errors: list[str] = []

    header = find_cell(ws, "Trecho do Circuito")
    col_secao_cell = find_cell(ws, "Seção")
    if not header or not col_secao_cell:
        return ["Cabecalho principal incompleto para validar entradas da planilha"]

    row0 = header.row
    col_trecho = header.column
    col_secao = col_secao_cell.column
    col_comp = col_secao + 4
    col_carga_kva = col_secao + 36

    trecho_regex = re.compile(r"^(TR(?:AF\.?\d*)?|P-\d+|RAMAL)$", flags=re.IGNORECASE)
    for r in range(row0 + 3, min(ws.max_row, row0 + 90)):
        trecho = str(ws.cell(row=r, column=col_trecho).value or "").strip()
        if not trecho_regex.match(trecho):
            continue

        secao = str(ws.cell(row=r, column=col_secao).value or "").strip()
        vao_m = to_float(ws.cell(row=r, column=col_comp).value, 0.0)
        carga_kva = to_float(ws.cell(row=r, column=col_carga_kva).value, 0.0)

        if vao_m > 0 and not secao:
            errors.append(f"{ws.title} {trecho}: vao preenchido sem condutor")
        if vao_m > 0 and carga_kva <= 0:
            errors.append(f"{ws.title} {trecho}: vao preenchido sem demanda/carga")

    return errors


def _audit_file(workbook_path: Path) -> dict[str, str]:
    result = {
        "Arquivo": str(workbook_path),
        "Status Paridade": "ERRO_PROCESSAMENTO",
        "Erros Humanos Detectados": "",
        "Queda Max Python": "",
        "Queda Max Excel": "",
        "Divergência": "",
    }

    wb = None
    temp_dir = None
    try:
        wb, temp_dir = _load_workbook_with_copy(workbook_path)
        sheets = list(_iter_cqt_worksheets(wb))
        if not sheets:
            raise RuntimeError("Nenhuma aba de CQT encontrada")

        parity_failures: list[str] = []
        human_errors: list[str] = []
        py_quedas: list[float] = []
        excel_quedas: list[float] = []
        rel_diffs: list[float] = []

        for ws in sheets:
            parsed = extract_cqt(ws)
            analise = CQTAnalise.model_validate(parsed.payload)

            py_qdt = float(analise.centro_carga.queda_total_percent)
            ex_qdt = float(parsed.expected_queda_total_percent)
            py_quedas.append(py_qdt)
            excel_quedas.append(ex_qdt)

            if abs(ex_qdt) > 1e-12:
                rel_diffs.append(abs(py_qdt - ex_qdt) / abs(ex_qdt) * 100.0)
            else:
                rel_diffs.append(0.0)

            if not _is_close(py_qdt, ex_qdt):
                parity_failures.append(ws.title)

            human_errors.extend(_collect_human_errors_from_sheet(ws))

            if py_qdt > QDT_LIMIT_PERCENT:
                human_errors.append(
                    f"{ws.title}: queda acumulada Python {py_qdt:.4f}% acima do limite regulatorio de {QDT_LIMIT_PERCENT:.2f}%"
                )
            if ex_qdt > QDT_LIMIT_PERCENT:
                human_errors.append(
                    f"{ws.title}: queda acumulada Excel {ex_qdt:.4f}% acima do limite regulatorio de {QDT_LIMIT_PERCENT:.2f}%"
                )

            for trecho in analise.centro_carga.trechos:
                if not trecho.dentro_do_limite_qdt:
                    human_errors.append(
                        f"{ws.title} {trecho.nome}: qdt do trecho {trecho.queda_tensao_percent:.4f}% acima do limite {trecho.limite_qdt_percent:.2f}%"
                    )

            esforco_excel = _extract_esforco_resultante(ws)
            if esforco_excel is None:
                human_errors.append(f"{ws.title}: esforco resultante nao identificado")

        max_py = max(py_quedas) if py_quedas else 0.0
        max_ex = max(excel_quedas) if excel_quedas else 0.0
        max_diff = max(rel_diffs) if rel_diffs else 0.0

        result["Status Paridade"] = (
            "OK" if not parity_failures else f"DIVERGENTE ({', '.join(parity_failures)})"
        )
        dedup_errors = list(dict.fromkeys(human_errors))
        result["Erros Humanos Detectados"] = "; ".join(dedup_errors) if dedup_errors else "Nenhum"
        result["Queda Max Python"] = f"{max_py:.6f}"
        result["Queda Max Excel"] = f"{max_ex:.6f}"
        result["Divergência"] = f"{max_diff:.6f}%"

    except ValidationError as exc:
        result["Erros Humanos Detectados"] = f"Falha de validacao: {exc}"
    except BaseException as exc:
        result["Erros Humanos Detectados"] = f"Falha de processamento: {exc}"
    finally:
        if wb is not None:
            wb.close()
        if temp_dir is not None:
            try:
                temp_dir.cleanup()
            except OSError:
                pass

    return result


def main() -> int:
    files = sorted(TARGET_DIR.rglob("QDT*.xlsm"))
    if not files:
        print(f"Nenhum arquivo QDT*.xlsm encontrado em: {TARGET_DIR}")
        return 1

    rows = [_audit_file(path) for path in files]

    with REPORT_PATH.open("w", newline="", encoding="utf-8") as csvfile:
        writer = csv.DictWriter(
            csvfile,
            fieldnames=[
                "Arquivo",
                "Status Paridade",
                "Erros Humanos Detectados",
                "Queda Max Python",
                "Queda Max Excel",
                "Divergência",
            ],
        )
        writer.writeheader()
        writer.writerows(rows)

    total = len(rows)
    parity_ok = sum(1 for row in rows if row["Status Paridade"] == "OK")
    parity_div = sum(1 for row in rows if row["Status Paridade"].startswith("DIVERGENTE"))
    proc_error = sum(1 for row in rows if row["Status Paridade"] == "ERRO_PROCESSAMENTO")
    human_error = sum(
        1
        for row in rows
        if row["Erros Humanos Detectados"] and row["Erros Humanos Detectados"] != "Nenhum"
    )

    print("Auditoria CQT legado concluida")
    print(f"- Diretorio auditado: {TARGET_DIR}")
    print(f"- Arquivos processados: {total}")
    print(f"- Paridade OK: {parity_ok}")
    print(f"- Paridade divergente: {parity_div}")
    print(f"- Erro de processamento: {proc_error}")
    print(f"- Com erros humanos detectados: {human_error}")
    print(f"- Relatorio: {REPORT_PATH}")

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
