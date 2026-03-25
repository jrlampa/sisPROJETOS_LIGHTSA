"""Rotas HTTP para execução do motor CQT."""

from __future__ import annotations

import io
import logging
import math
from uuid import UUID

import openpyxl
from fastapi import APIRouter, Depends, File, HTTPException, UploadFile, status
from pydantic import ValidationError

from apps.api.core.deps import get_session_factory
from apps.api.core.security import require_write_access
from apps.api.schemas.cqt_schemas import CQTAnaliseRequest, CQTAnaliseResponse
from packages.application.use_cases.cqt_use_cases import ExecutarAnaliseCQTUseCase
from packages.domain.tests.excel_parity_support import extract_cqt, find_cell, norm_text, to_float
from packages.infrastructure.excel_importer.parser import (
    CQTParserError,
)

router = APIRouter(tags=["cqt"])
logger = logging.getLogger(__name__)


def _iter_cqt_worksheets(wb):
    for ws in wb.worksheets:
        has_trecho = find_cell(ws, "Trecho do Circuito") is not None
        has_queda = find_cell(ws, "Queda") is not None or find_cell(ws, "DV") is not None
        if has_trecho and has_queda:
            yield ws


def _lado_from_title(title: str) -> str | None:
    normalized = norm_text(title)
    if "ESQUER" in normalized or "LADO2" in normalized:
        return "esquerdo"
    if "DIREIT" in normalized or "LADO1" in normalized:
        return "direito"
    return None


def _to_num_str(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float):
        if math.isnan(value):
            return ""
        if value.is_integer():
            return str(int(value))
    return str(value)


def _build_estado_from_sheet(ws) -> dict:
    parsed = extract_cqt(ws)
    trecho_header = find_cell(ws, "Trecho do Circuito")
    secao_header = find_cell(ws, "Seção")

    if trecho_header is None or secao_header is None:
        raise AssertionError("Nao foi possivel mapear cabecalhos essenciais da planilha CQT.")

    col_trecho = trecho_header.column
    row0 = trecho_header.row
    col_secao = secao_header.column
    col_fases = (
        find_cell(ws, "Nº de fases do trecho")
        or find_cell(ws, "No de fases do trecho")
        or secao_header
    ).column
    col_comp = col_secao + 4
    col_carga_kva = col_secao + 36
    col_qdt_trecho = col_secao + 39
    col_qdt_acum = col_secao + 40

    trechos_raw = parsed.payload["centro_carga"]["trechos"]
    linhas = []
    for idx, trecho in enumerate(trechos_raw, start=1):
        row = row0 + 2 + idx
        nome_trecho = str(ws.cell(row=row, column=col_trecho).value or trecho.get("nome") or "")
        nome_upper = norm_text(nome_trecho)
        poste = ""
        if nome_upper.startswith("P-"):
            poste = nome_trecho.split("-", 1)[-1].strip()

        linhas.append(
            {
                "id": idx,
                "poste": poste,
                "vao_m": _to_num_str(
                    to_float(
                        ws.cell(row=row, column=col_comp).value, trecho.get("comprimento_m", 0.0)
                    )
                ),
                "corrente_a": _to_num_str(trecho.get("corrente_a", 0.0)),
                "esforco_dan": _to_num_str(
                    to_float(ws.cell(row=row, column=col_carga_kva).value, 0.0)
                ),
                "tipo_cabo": _to_num_str(trecho.get("condutor", {}).get("nome", "")),
                "fases": _to_num_str(
                    to_float(ws.cell(row=row, column=col_fases).value, trecho.get("fases", 3))
                ),
                "queda_tensao_trecho": _to_num_str(
                    to_float(
                        ws.cell(row=row, column=col_qdt_trecho).value,
                        trecho.get("queda_tensao_percent_legacy", 0.0),
                    )
                ),
                "queda_tensao_acumulada": _to_num_str(
                    to_float(ws.cell(row=row, column=col_qdt_acum).value, 0.0)
                ),
                "fim_linha": "Sim" if idx == len(trechos_raw) else "Nao",
            }
        )

    cabecalho = {
        "nomeProjeto": "",
        "projetista": "",
        "data": "",
        "localidade": "",
        "condutores": _to_num_str(linhas[0]["tipo_cabo"] if linhas else ""),
        "demanda": _to_num_str(
            max((to_float(linha["corrente_a"], 0.0) for linha in linhas), default=0.0)
        ),
        "trafoKva": _to_num_str(
            parsed.payload.get("centro_carga", {})
            .get("transformador", {})
            .get("potencia_nominal_kva", "")
        ),
        "tensao": "220",
        "fatorPotencia": "0.92",
        "observacoes": "",
    }
    return {"cabecalho": cabecalho, "trechos": linhas}


@router.post(
    "/projetos/{projeto_id}/cqt",
    response_model=CQTAnaliseResponse,
    status_code=status.HTTP_201_CREATED,
)
def executar_analise_cqt(
    projeto_id: UUID,
    payload: CQTAnaliseRequest,
    session_factory=Depends(get_session_factory),
    _: object = Depends(require_write_access),
) -> CQTAnaliseResponse:
    """Executa análise CQT, persiste os resultados e retorna cálculos consolidados."""

    use_case = ExecutarAnaliseCQTUseCase(session_factory=session_factory)
    try:
        analise = use_case.executar(
            projeto_id=projeto_id,
            dados_brutos_cqt=payload.model_dump(mode="json", exclude_none=True),
        )
    except ValidationError as exc:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=str(exc)) from exc
    except ValueError as exc:
        message = str(exc)
        code = (
            status.HTTP_404_NOT_FOUND
            if "nao encontrado" in message.lower()
            else status.HTTP_400_BAD_REQUEST
        )
        raise HTTPException(status_code=code, detail=message) from exc

    logger.info("Analise CQT concluida com sucesso: projeto_id=%s", projeto_id)
    return CQTAnaliseResponse.from_domain(analise)


@router.post(
    "/cqt/importar-excel",
    status_code=status.HTTP_200_OK,
)
async def importar_excel_cqt(
    file: UploadFile = File(...),
    _: object = Depends(require_write_access),
) -> dict:
    """Importa planilha CQT legada e retorna estrutura pronta para hidratar o frontend."""

    filename = (file.filename or "").lower()
    if not (filename.endswith(".xlsm") or filename.endswith(".xlsx")):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Formato invalido. Envie um arquivo .xlsm ou .xlsx.",
        )

    raw = await file.read()
    if not raw:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Arquivo vazio.",
        )

    wb = None
    try:
        wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=True, read_only=False)
        sheets = list(_iter_cqt_worksheets(wb))
        if not sheets:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Nao foi encontrada aba CQT no arquivo enviado.",
            )

        estados: dict[str, dict] = {}
        fallback = []
        for ws in sheets:
            estado = _build_estado_from_sheet(ws)
            lado = _lado_from_title(ws.title)
            if lado is None:
                fallback.append(estado)
                continue
            estados[lado] = estado

        if "esquerdo" not in estados and fallback:
            estados["esquerdo"] = fallback.pop(0)
        if "direito" not in estados and fallback:
            estados["direito"] = fallback.pop(0)

        if "esquerdo" not in estados or "direito" not in estados:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Nao foi possivel identificar os lados esquerdo e direito na planilha CQT.",
            )

        return {
            "esquerdo": estados["esquerdo"],
            "direito": estados["direito"],
        }
    except CQTParserError as exc:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Erro de parsing na planilha: {exc}",
        ) from exc
    except HTTPException:
        raise
    except Exception as exc:
        logger.exception("Falha inesperada ao importar planilha CQT.")
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Falha ao ler a planilha: {exc}",
        ) from exc
    finally:
        try:
            if wb is not None:
                wb.close()
        except Exception:
            pass
